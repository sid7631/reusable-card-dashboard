"""Ingest uploaded CSV / Excel / JSON files into a per-file DuckDB database."""

import json
from pathlib import Path
from typing import Any

import pyarrow as pa
from openpyxl import load_workbook

from app.connectors import duck
from app.connectors.common import QueryError

FORMATS = {".csv": "csv", ".tsv": "csv", ".xlsx": "xlsx", ".xlsm": "xlsx", ".json": "json"}


def detect_format(filename: str) -> str:
    fmt = FORMATS.get(Path(filename).suffix.lower())
    if fmt is None:
        raise QueryError("Unsupported file type — upload CSV, Excel (.xlsx) or JSON.")
    return fmt


def records_at(payload: Any, path: str = "") -> list[dict[str, Any]]:
    """Find the record list: an explicit dot path, the root array, or the only array in the root object."""
    node = payload
    for part in filter(None, path.split(".")):
        if not isinstance(node, dict) or part not in node:
            raise QueryError(f"Records path “{path}” not found in the response.")
        node = node[part]
    if isinstance(node, dict) and not path:
        lists = [v for v in node.values() if isinstance(v, list)]
        node = lists[0] if len(lists) == 1 else [node]
    if isinstance(node, dict):
        node = [node]
    if not isinstance(node, list):
        raise QueryError("Expected a list of records.")
    return [flatten(r) if isinstance(r, dict) else {"value": r} for r in node]


def flatten(record: dict[str, Any], prefix: str = "") -> dict[str, Any]:
    """One table row per record: nested objects become prefix_key columns, lists become JSON text."""
    out: dict[str, Any] = {}
    for k, v in record.items():
        key = duck.safe_table_name(f"{prefix}{k}", "col")
        if isinstance(v, dict):
            out.update(flatten(v, f"{key}_"))
        elif isinstance(v, list):
            out[key] = json.dumps(v, default=str)
        else:
            out[key] = v
    return out


def arrow_table(rows: list[dict[str, Any]]) -> pa.Table:
    """Build an Arrow table, falling back to text for columns whose values have mixed types."""
    if not rows:
        raise QueryError("No records found.")
    columns: dict[str, list[Any]] = {}
    for i, row in enumerate(rows):
        for k in row:
            columns.setdefault(k, [None] * i)
        for k, values in columns.items():
            values.append(row.get(k))
    arrays = {}
    for name, values in columns.items():
        kinds = {type(v) for v in values if v is not None}
        if kinds <= {int, float} or len(kinds) <= 1:
            try:
                arrays[name] = pa.array(values)
                continue
            except (pa.ArrowInvalid, pa.ArrowTypeError, OverflowError):
                pass
        arrays[name] = pa.array([None if v is None else str(v) for v in values], type=pa.string())
    return pa.table(arrays)


def _excel_sheets(path: Path) -> dict[str, pa.Table]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = {}
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                continue
            names = [duck.safe_table_name(str(h), f"col_{i + 1}") if h is not None else f"col_{i + 1}"
                     for i, h in enumerate(header)]
            records = [dict(zip(names, row, strict=False)) for row in rows if any(v is not None for v in row)]
            if records:
                sheets[duck.safe_table_name(ws.title, "sheet")] = arrow_table(records)
        return sheets
    finally:
        wb.close()


def ingest(raw_path: Path, fmt: str, db_path: Path, table: str) -> list[dict[str, Any]]:
    """Load the file into `db_path` (one table, or one per Excel sheet) and describe the result."""
    def build(con):
        if fmt == "csv":
            con.execute(f"CREATE TABLE \"{table}\" AS SELECT * FROM read_csv_auto('{raw_path.as_posix()}', sample_size=-1)")
        elif fmt == "json":
            with raw_path.open() as fh:
                payload = json.load(fh)
            con.register("_upload", arrow_table(records_at(payload)))
            con.execute(f'CREATE TABLE "{table}" AS SELECT * FROM _upload')
        else:
            sheets = _excel_sheets(raw_path)
            if not sheets:
                raise QueryError("The workbook has no data.")
            for name, data in sheets.items():
                con.register("_sheet", data)
                con.execute(f'CREATE TABLE "{name}" AS SELECT * FROM _sheet')
                con.unregister("_sheet")

    try:
        duck.build_database(db_path, build)
    except QueryError:
        raise
    except (json.JSONDecodeError, UnicodeDecodeError) as exc:
        raise QueryError(f"Could not read the file: {exc}") from exc
    except Exception as exc:  # duckdb/openpyxl parse errors
        raise QueryError(f"Could not read the file: {str(exc).splitlines()[0]}") from exc
    return duck.describe(db_path)
